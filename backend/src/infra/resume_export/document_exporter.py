from dataclasses import dataclass
from datetime import date
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.document import Document as WordDocument
from docx.shared import Inches
from docx.text.paragraph import Paragraph as WordParagraph
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer
from reportlab.platypus.flowables import Flowable

from core.i18n.enums import LanguageEnum
from core.resumes.enums import ResumeCurrentStatusEnum, ResumeExportFormatEnum
from core.resumes.exporters import ResumeDocumentExporter
from core.resumes.schemas import (
    ResumeAdditionalSection,
    ResumeCertificationItem,
    ResumeContent,
    ResumeEducationItem,
    ResumeExperienceItem,
    ResumeExport,
    ResumeExportParams,
    ResumeProjectItem,
)


@dataclass(frozen=True, slots=True, kw_only=True)
class ResumeExportLabels:
    summary: str
    skills: str
    experience: str
    projects: str
    education: str
    languages: str
    certifications: str
    present: str
    technologies: str


@dataclass(frozen=True, slots=True, kw_only=True)
class ResumePdfStyles:
    title: ParagraphStyle
    subtitle: ParagraphStyle
    contact: ParagraphStyle
    section: ParagraphStyle
    item_title: ParagraphStyle
    body: ParagraphStyle
    bullet: ParagraphStyle


@dataclass(frozen=True, slots=True, kw_only=True)
class ResumeDocumentExporterImpl(ResumeDocumentExporter):
    font_regular_path: Path
    font_bold_path: Path
    font_regular_name: str
    font_bold_name: str

    def export_resume(self, *, params: ResumeExportParams) -> ResumeExport:
        if params.format == ResumeExportFormatEnum.PDF:
            return ResumeExport(
                format=params.format,
                content=self._export_pdf(params=params),
            )
        if params.format == ResumeExportFormatEnum.DOCX:
            return ResumeExport(
                format=params.format,
                content=self._export_docx(params=params),
            )
        message = f"Unsupported resume export format: {params.format}"
        raise ValueError(message)

    def _export_pdf(self, *, params: ResumeExportParams) -> bytes:
        self._register_pdf_fonts()
        styles = self._build_pdf_styles()
        labels = self._labels_for_language(language=params.language)
        output = BytesIO()
        document = SimpleDocTemplate(
            output,
            pagesize=A4,
            rightMargin=18 * mm,
            leftMargin=18 * mm,
            topMargin=16 * mm,
            bottomMargin=16 * mm,
            title=params.title,
            author=params.content.profile.full_name,
        )
        story: list[Flowable] = []
        self._append_pdf_profile(
            story=story,
            title=params.title,
            content=params.content,
            styles=styles,
        )
        self._append_pdf_summary(
            story=story,
            content=params.content,
            labels=labels,
            styles=styles,
        )
        self._append_pdf_skills(
            story=story,
            content=params.content,
            labels=labels,
            styles=styles,
        )
        self._append_pdf_experience(
            story=story,
            content=params.content,
            labels=labels,
            styles=styles,
        )
        self._append_pdf_education(
            story=story,
            content=params.content,
            labels=labels,
            styles=styles,
        )
        self._append_pdf_languages(
            story=story,
            content=params.content,
            labels=labels,
            styles=styles,
        )
        self._append_pdf_certifications(
            story=story,
            content=params.content,
            labels=labels,
            styles=styles,
        )
        self._append_pdf_additional_sections(
            story=story,
            sections=params.content.additional_sections,
            styles=styles,
        )
        document.build(story)
        return output.getvalue()

    def _export_docx(self, *, params: ResumeExportParams) -> bytes:
        document = Document()
        self._configure_word_document(document=document)
        labels = self._labels_for_language(language=params.language)
        self._append_word_profile(
            document=document,
            title=params.title,
            content=params.content,
        )
        self._append_word_summary(document=document, content=params.content, labels=labels)
        self._append_word_skills(document=document, content=params.content, labels=labels)
        self._append_word_experience(document=document, content=params.content, labels=labels)
        self._append_word_education(document=document, content=params.content, labels=labels)
        self._append_word_languages(document=document, content=params.content, labels=labels)
        self._append_word_certifications(document=document, content=params.content, labels=labels)
        self._append_word_additional_sections(
            document=document,
            sections=params.content.additional_sections,
        )
        output = BytesIO()
        document.save(output)
        return output.getvalue()

    def _register_pdf_fonts(self) -> None:
        pdfmetrics.registerFont(TTFont(self.font_regular_name, self.font_regular_path))
        pdfmetrics.registerFont(TTFont(self.font_bold_name, self.font_bold_path))

    def _build_pdf_styles(self) -> ResumePdfStyles:
        sample_styles = getSampleStyleSheet()
        return ResumePdfStyles(
            title=ParagraphStyle(
                name="ResumeTitle",
                parent=sample_styles["Title"],
                fontName=self.font_bold_name,
                fontSize=18,
                leading=22,
                spaceAfter=4,
            ),
            subtitle=ParagraphStyle(
                name="ResumeSubtitle",
                parent=sample_styles["Normal"],
                fontName=self.font_regular_name,
                fontSize=11,
                leading=15,
                spaceAfter=8,
            ),
            contact=ParagraphStyle(
                name="ResumeContact",
                parent=sample_styles["Normal"],
                fontName=self.font_regular_name,
                fontSize=9,
                leading=12,
                spaceAfter=8,
            ),
            section=ParagraphStyle(
                name="ResumeSection",
                parent=sample_styles["Heading2"],
                fontName=self.font_bold_name,
                fontSize=12,
                leading=15,
                spaceBefore=8,
                spaceAfter=4,
            ),
            item_title=ParagraphStyle(
                name="ResumeItemTitle",
                parent=sample_styles["Normal"],
                fontName=self.font_bold_name,
                fontSize=10,
                leading=13,
                spaceBefore=2,
                spaceAfter=2,
            ),
            body=ParagraphStyle(
                name="ResumeBody",
                parent=sample_styles["Normal"],
                fontName=self.font_regular_name,
                fontSize=9,
                leading=12,
                spaceAfter=3,
            ),
            bullet=ParagraphStyle(
                name="ResumeBullet",
                parent=sample_styles["Normal"],
                fontName=self.font_regular_name,
                fontSize=9,
                leading=12,
            ),
        )

    def _configure_word_document(self, *, document: WordDocument) -> None:
        section = document.sections[0]
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)
        for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet"):
            style = document.styles[style_name]
            style.font.name = self.font_regular_name

    def _labels_for_language(self, *, language: LanguageEnum) -> ResumeExportLabels:
        if language == LanguageEnum.RU:
            return ResumeExportLabels(
                summary="Профиль",
                skills="Навыки",
                experience="Опыт",
                projects="Проекты",
                education="Образование",
                languages="Языки",
                certifications="Сертификаты",
                present="по настоящее время",
                technologies="Технологии",
            )
        return ResumeExportLabels(
            summary="Profile",
            skills="Skills",
            experience="Experience",
            projects="Projects",
            education="Education",
            languages="Languages",
            certifications="Certifications",
            present="Present",
            technologies="Technologies",
        )

    def _append_pdf_profile(
        self,
        *,
        story: list[Flowable],
        title: str,
        content: ResumeContent,
        styles: ResumePdfStyles,
    ) -> None:
        heading = content.profile.full_name or title
        self._append_pdf_text(story=story, text=heading, style=styles.title)
        self._append_pdf_text(story=story, text=content.profile.role, style=styles.subtitle)
        contact_parts = [
            content.profile.location,
            content.profile.email,
            content.profile.phone,
            content.profile.website_url,
            content.profile.linkedin_url,
            content.profile.github_url,
            content.profile.telegram,
        ]
        self._append_pdf_text(
            story=story,
            text=" | ".join(part for part in contact_parts if part),
            style=styles.contact,
        )

    def _append_pdf_summary(
        self,
        *,
        story: list[Flowable],
        content: ResumeContent,
        labels: ResumeExportLabels,
        styles: ResumePdfStyles,
    ) -> None:
        if not content.summary.text:
            return
        self._append_pdf_text(story=story, text=labels.summary, style=styles.section)
        self._append_pdf_text(story=story, text=content.summary.text, style=styles.body)

    def _append_pdf_skills(
        self,
        *,
        story: list[Flowable],
        content: ResumeContent,
        labels: ResumeExportLabels,
        styles: ResumePdfStyles,
    ) -> None:
        if not content.skills:
            return
        self._append_pdf_text(story=story, text=labels.skills, style=styles.section)
        for skill_group in content.skills:
            text_parts = [skill_group.category, ", ".join(skill_group.items)]
            self._append_pdf_text(
                story=story,
                text=": ".join(part for part in text_parts if part),
                style=styles.body,
            )

    def _append_pdf_experience(
        self,
        *,
        story: list[Flowable],
        content: ResumeContent,
        labels: ResumeExportLabels,
        styles: ResumePdfStyles,
    ) -> None:
        if not content.experience:
            return
        self._append_pdf_text(story=story, text=labels.experience, style=styles.section)
        for item in content.experience:
            self._append_pdf_experience_item(
                story=story,
                item=item,
                labels=labels,
                styles=styles,
            )

    def _append_pdf_experience_item(
        self,
        *,
        story: list[Flowable],
        item: ResumeExperienceItem,
        labels: ResumeExportLabels,
        styles: ResumePdfStyles,
    ) -> None:
        title_parts = [item.position, item.company]
        self._append_pdf_text(
            story=story,
            text=" - ".join(part for part in title_parts if part),
            style=styles.item_title,
        )
        meta_parts = [
            item.location,
            self._format_date_range(
                start_date=item.start_date,
                end_date=item.end_date,
                current_status=item.current_status,
                labels=labels,
            ),
        ]
        self._append_pdf_text(
            story=story,
            text=" | ".join(part for part in meta_parts if part),
            style=styles.body,
        )
        self._append_pdf_text(story=story, text=item.summary, style=styles.body)
        self._append_pdf_bullets(story=story, items=item.highlights, styles=styles)
        self._append_pdf_technologies(
            story=story,
            technologies=item.technologies,
            labels=labels,
            styles=styles,
        )
        self._append_pdf_projects(
            story=story,
            projects=item.projects,
            labels=labels,
            styles=styles,
        )

    def _append_pdf_projects(
        self,
        *,
        story: list[Flowable],
        projects: list[ResumeProjectItem],
        labels: ResumeExportLabels,
        styles: ResumePdfStyles,
    ) -> None:
        if not projects:
            return
        self._append_pdf_text(story=story, text=labels.projects, style=styles.item_title)
        for project in projects:
            title_parts = [project.name, project.role]
            self._append_pdf_text(
                story=story,
                text=" - ".join(part for part in title_parts if part),
                style=styles.body,
            )
            self._append_pdf_text(story=story, text=project.description, style=styles.body)
            self._append_pdf_bullets(story=story, items=project.highlights, styles=styles)
            self._append_pdf_technologies(
                story=story,
                technologies=project.technologies,
                labels=labels,
                styles=styles,
            )
            self._append_pdf_text(story=story, text=project.url, style=styles.body)

    def _append_pdf_education(
        self,
        *,
        story: list[Flowable],
        content: ResumeContent,
        labels: ResumeExportLabels,
        styles: ResumePdfStyles,
    ) -> None:
        if not content.education:
            return
        self._append_pdf_text(story=story, text=labels.education, style=styles.section)
        for item in content.education:
            self._append_pdf_education_item(
                story=story,
                item=item,
                labels=labels,
                styles=styles,
            )

    def _append_pdf_education_item(
        self,
        *,
        story: list[Flowable],
        item: ResumeEducationItem,
        labels: ResumeExportLabels,
        styles: ResumePdfStyles,
    ) -> None:
        title_parts = [item.institution, item.degree, item.field]
        self._append_pdf_text(
            story=story,
            text=" - ".join(part for part in title_parts if part),
            style=styles.item_title,
        )
        meta_parts = [
            item.location,
            self._format_date_range(
                start_date=item.start_date,
                end_date=item.end_date,
                current_status=ResumeCurrentStatusEnum.NOT_SET,
                labels=labels,
            ),
        ]
        self._append_pdf_text(
            story=story,
            text=" | ".join(part for part in meta_parts if part),
            style=styles.body,
        )
        self._append_pdf_text(story=story, text=item.description, style=styles.body)

    def _append_pdf_languages(
        self,
        *,
        story: list[Flowable],
        content: ResumeContent,
        labels: ResumeExportLabels,
        styles: ResumePdfStyles,
    ) -> None:
        if not content.languages:
            return
        self._append_pdf_text(story=story, text=labels.languages, style=styles.section)
        for item in content.languages:
            self._append_pdf_text(
                story=story,
                text=" - ".join(part for part in (item.name, item.proficiency) if part),
                style=styles.body,
            )

    def _append_pdf_certifications(
        self,
        *,
        story: list[Flowable],
        content: ResumeContent,
        labels: ResumeExportLabels,
        styles: ResumePdfStyles,
    ) -> None:
        if not content.certifications:
            return
        self._append_pdf_text(story=story, text=labels.certifications, style=styles.section)
        for item in content.certifications:
            self._append_pdf_certification_item(
                story=story,
                item=item,
                styles=styles,
            )

    def _append_pdf_certification_item(
        self,
        *,
        story: list[Flowable],
        item: ResumeCertificationItem,
        styles: ResumePdfStyles,
    ) -> None:
        title_parts = [item.name, item.issuer]
        self._append_pdf_text(
            story=story,
            text=" - ".join(part for part in title_parts if part),
            style=styles.item_title,
        )
        date_parts = [self._format_date(item.issued_on), self._format_date(item.expires_on)]
        self._append_pdf_text(
            story=story,
            text=" - ".join(part for part in date_parts if part),
            style=styles.body,
        )
        self._append_pdf_text(story=story, text=item.credential_url, style=styles.body)

    def _append_pdf_additional_sections(
        self,
        *,
        story: list[Flowable],
        sections: list[ResumeAdditionalSection],
        styles: ResumePdfStyles,
    ) -> None:
        for section in sections:
            if not section.items:
                continue
            self._append_pdf_text(story=story, text=section.title, style=styles.section)
            for item in section.items:
                self._append_pdf_text(story=story, text=item.title, style=styles.item_title)
                self._append_pdf_text(story=story, text=item.description, style=styles.body)
                self._append_pdf_text(story=story, text=item.url, style=styles.body)

    def _append_pdf_technologies(
        self,
        *,
        story: list[Flowable],
        technologies: list[str],
        labels: ResumeExportLabels,
        styles: ResumePdfStyles,
    ) -> None:
        if not technologies:
            return
        self._append_pdf_text(
            story=story,
            text=f"{labels.technologies}: {', '.join(technologies)}",
            style=styles.body,
        )

    def _append_pdf_text(
        self,
        *,
        story: list[Flowable],
        text: str,
        style: ParagraphStyle,
    ) -> None:
        if not text:
            return
        story.append(Paragraph(escape(text), style))

    def _append_pdf_bullets(
        self,
        *,
        story: list[Flowable],
        items: list[str],
        styles: ResumePdfStyles,
    ) -> None:
        bullets = [
            ListItem(Paragraph(escape(item), styles.bullet), leftIndent=8) for item in items if item
        ]
        if not bullets:
            return
        story.append(ListFlowable(bullets, bulletType="bullet", leftIndent=14))
        story.append(Spacer(1, 2))

    def _append_word_profile(
        self,
        *,
        document: WordDocument,
        title: str,
        content: ResumeContent,
    ) -> None:
        heading = content.profile.full_name or title
        self._add_word_heading(document=document, text=heading, level=0)
        self._add_word_paragraph(document=document, text=content.profile.role)
        contact_parts = [
            content.profile.location,
            content.profile.email,
            content.profile.phone,
            content.profile.website_url,
            content.profile.linkedin_url,
            content.profile.github_url,
            content.profile.telegram,
        ]
        self._add_word_paragraph(
            document=document,
            text=" | ".join(part for part in contact_parts if part),
        )

    def _append_word_summary(
        self,
        *,
        document: WordDocument,
        content: ResumeContent,
        labels: ResumeExportLabels,
    ) -> None:
        if not content.summary.text:
            return
        self._add_word_heading(document=document, text=labels.summary, level=1)
        self._add_word_paragraph(document=document, text=content.summary.text)

    def _append_word_skills(
        self,
        *,
        document: WordDocument,
        content: ResumeContent,
        labels: ResumeExportLabels,
    ) -> None:
        if not content.skills:
            return
        self._add_word_heading(document=document, text=labels.skills, level=1)
        for skill_group in content.skills:
            text_parts = [skill_group.category, ", ".join(skill_group.items)]
            self._add_word_paragraph(
                document=document,
                text=": ".join(part for part in text_parts if part),
            )

    def _append_word_experience(
        self,
        *,
        document: WordDocument,
        content: ResumeContent,
        labels: ResumeExportLabels,
    ) -> None:
        if not content.experience:
            return
        self._add_word_heading(document=document, text=labels.experience, level=1)
        for item in content.experience:
            title_parts = [item.position, item.company]
            self._add_word_heading(
                document=document,
                text=" - ".join(part for part in title_parts if part),
                level=2,
            )
            meta_parts = [
                item.location,
                self._format_date_range(
                    start_date=item.start_date,
                    end_date=item.end_date,
                    current_status=item.current_status,
                    labels=labels,
                ),
            ]
            self._add_word_paragraph(
                document=document,
                text=" | ".join(part for part in meta_parts if part),
            )
            self._add_word_paragraph(document=document, text=item.summary)
            for highlight in item.highlights:
                self._add_word_bullet(document=document, text=highlight)
            self._add_word_technologies(
                document=document,
                technologies=item.technologies,
                labels=labels,
            )
            self._append_word_projects(
                document=document,
                projects=item.projects,
                labels=labels,
            )

    def _append_word_projects(
        self,
        *,
        document: WordDocument,
        projects: list[ResumeProjectItem],
        labels: ResumeExportLabels,
    ) -> None:
        if not projects:
            return
        self._add_word_heading(document=document, text=labels.projects, level=2)
        for project in projects:
            title_parts = [project.name, project.role]
            self._add_word_paragraph(
                document=document,
                text=" - ".join(part for part in title_parts if part),
            )
            self._add_word_paragraph(document=document, text=project.description)
            for highlight in project.highlights:
                self._add_word_bullet(document=document, text=highlight)
            self._add_word_technologies(
                document=document,
                technologies=project.technologies,
                labels=labels,
            )
            self._add_word_paragraph(document=document, text=project.url)

    def _append_word_education(
        self,
        *,
        document: WordDocument,
        content: ResumeContent,
        labels: ResumeExportLabels,
    ) -> None:
        if not content.education:
            return
        self._add_word_heading(document=document, text=labels.education, level=1)
        for item in content.education:
            title_parts = [item.institution, item.degree, item.field]
            self._add_word_heading(
                document=document,
                text=" - ".join(part for part in title_parts if part),
                level=2,
            )
            meta_parts = [
                item.location,
                self._format_date_range(
                    start_date=item.start_date,
                    end_date=item.end_date,
                    current_status=ResumeCurrentStatusEnum.NOT_SET,
                    labels=labels,
                ),
            ]
            self._add_word_paragraph(
                document=document,
                text=" | ".join(part for part in meta_parts if part),
            )
            self._add_word_paragraph(document=document, text=item.description)

    def _append_word_languages(
        self,
        *,
        document: WordDocument,
        content: ResumeContent,
        labels: ResumeExportLabels,
    ) -> None:
        if not content.languages:
            return
        self._add_word_heading(document=document, text=labels.languages, level=1)
        for item in content.languages:
            self._add_word_paragraph(
                document=document,
                text=" - ".join(part for part in (item.name, item.proficiency) if part),
            )

    def _append_word_certifications(
        self,
        *,
        document: WordDocument,
        content: ResumeContent,
        labels: ResumeExportLabels,
    ) -> None:
        if not content.certifications:
            return
        self._add_word_heading(document=document, text=labels.certifications, level=1)
        for item in content.certifications:
            title_parts = [item.name, item.issuer]
            self._add_word_heading(
                document=document,
                text=" - ".join(part for part in title_parts if part),
                level=2,
            )
            date_parts = [self._format_date(item.issued_on), self._format_date(item.expires_on)]
            self._add_word_paragraph(
                document=document,
                text=" - ".join(part for part in date_parts if part),
            )
            self._add_word_paragraph(document=document, text=item.credential_url)

    def _append_word_additional_sections(
        self,
        *,
        document: WordDocument,
        sections: list[ResumeAdditionalSection],
    ) -> None:
        for section in sections:
            if not section.items:
                continue
            self._add_word_heading(document=document, text=section.title, level=1)
            for item in section.items:
                self._add_word_heading(document=document, text=item.title, level=2)
                self._add_word_paragraph(document=document, text=item.description)
                self._add_word_paragraph(document=document, text=item.url)

    def _add_word_technologies(
        self,
        *,
        document: WordDocument,
        technologies: list[str],
        labels: ResumeExportLabels,
    ) -> None:
        if not technologies:
            return
        self._add_word_paragraph(
            document=document,
            text=f"{labels.technologies}: {', '.join(technologies)}",
        )

    def _add_word_heading(self, *, document: WordDocument, text: str, level: int) -> None:
        if not text:
            return
        paragraph = document.add_heading(text, level=level)
        self._apply_word_paragraph_font(paragraph=paragraph, bold=True)

    def _add_word_paragraph(self, *, document: WordDocument, text: str) -> None:
        if not text:
            return
        paragraph = document.add_paragraph(text)
        self._apply_word_paragraph_font(paragraph=paragraph, bold=False)

    def _add_word_bullet(self, *, document: WordDocument, text: str) -> None:
        if not text:
            return
        paragraph = document.add_paragraph(text, style="List Bullet")
        self._apply_word_paragraph_font(paragraph=paragraph, bold=False)

    def _apply_word_paragraph_font(self, *, paragraph: WordParagraph, bold: bool) -> None:
        for run in paragraph.runs:
            run.font.name = self.font_bold_name if bold else self.font_regular_name

    def _format_date_range(
        self,
        *,
        start_date: date | None,
        end_date: date | None,
        current_status: ResumeCurrentStatusEnum,
        labels: ResumeExportLabels,
    ) -> str:
        start = self._format_date(start_date)
        end = labels.present if current_status == ResumeCurrentStatusEnum.CURRENT else ""
        if not end:
            end = self._format_date(end_date)
        if start and end:
            return f"{start} - {end}"
        return start or end

    def _format_date(self, value: date | None) -> str:
        if value is None:
            return ""
        return value.isoformat()
